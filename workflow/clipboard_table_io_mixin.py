# -*- coding: utf-8 -*-
"""Main ClipboardTableApp SQLite and xlsx I/O helpers."""

import os
import re
import zipfile
from datetime import datetime
from xml.sax.saxutils import escape

from tkinter import filedialog, messagebox

from core.text_utils import (
    make_sql_columns as core_make_sql_columns,
    quote_ident as core_quote_ident,
    sanitize_sql_name as core_sanitize_sql_name,
)
from db import TableAccessManager


class ClipboardTableIoMixin:
    """Compatibility methods used by the main preview table and workflow windows."""

    def normalize_sheet_title(self, name):
        name = str(name or "导出数据").strip() or "导出数据"
        name = re.sub(r"[\\/*?:\[\]]", "_", name)
        return name[:31] or "导出数据"

    def column_letter(self, index):
        result = ""
        while index > 0:
            index, rem = divmod(index - 1, 26)
            result = chr(65 + rem) + result
        return result or "A"

    def calc_display_width(self, value):
        text = str(value or "")
        width = 0
        for ch in text:
            width += 2 if ord(ch) > 127 else 1
        return width

    def export_current_preview_to_xlsx(self, headers=None, rows=None, table_name=None, title="导出为 xlsx"):
        headers = list(self.headers if headers is None else headers)
        rows = [list(row) for row in (self.rows if rows is None else rows)]
        table_name = self.table_name_var.get() if table_name is None else table_name

        if not headers:
            messagebox.showwarning("提示", "当前没有可导出的表格字段。")
            return

        default_base = self.sanitize_sql_name(table_name, "导出数据")
        default_name = f"{default_base}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        path = filedialog.asksaveasfilename(
            title=title,
            defaultextension=".xlsx",
            initialfile=default_name,
            filetypes=[("Excel 工作簿", "*.xlsx"), ("所有文件", "*.*")]
        )

        if not path:
            return

        if not path.lower().endswith(".xlsx"):
            path += ".xlsx"

        try:
            try:
                self.export_xlsx_with_openpyxl(path, headers=headers, rows=rows, table_name=table_name)
                engine = "openpyxl"
            except ModuleNotFoundError:
                self.export_xlsx_minimal(path, headers=headers, rows=rows, table_name=table_name)
                engine = "内置简易导出器"

            self.info_var.set(f"导出成功：{path}")
            messagebox.showinfo(
                "导出成功",
                f"已导出当前预览数据。\n\n文件：{path}\n行数：{len(rows)}\n列数：{len(headers)}\n导出方式：{engine}"
            )
        except Exception as e:
            messagebox.showerror("导出失败", str(e))

    def export_xlsx_with_openpyxl(self, path, headers=None, rows=None, table_name=None):
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        headers = [str(h) for h in (self.headers if headers is None else headers)]
        rows = [list(row) for row in (self.rows if rows is None else rows)]
        table_name = self.table_name_var.get() if table_name is None else table_name

        wb = Workbook()
        ws = wb.active
        ws.title = self.normalize_sheet_title(table_name)

        ws.append(headers)

        for row in rows:
            fixed = list(row)
            if len(fixed) < len(headers):
                fixed += [""] * (len(headers) - len(fixed))
            if len(fixed) > len(headers):
                fixed = fixed[:len(headers)]
            ws.append(["" if value is None else str(value) for value in fixed])

        header_fill = PatternFill("solid", fgColor="D9EAF7")
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for row_cells in ws.iter_rows(min_row=2):
            for cell in row_cells:
                cell.alignment = Alignment(vertical="center")
                cell.border = border

        ws.freeze_panes = "A2"
        if headers:
            last_col = self.column_letter(len(headers))
            ws.auto_filter.ref = f"A1:{last_col}{max(len(rows) + 1, 1)}"

        for col_idx, header in enumerate(headers, start=1):
            max_width = self.calc_display_width(header)
            for row in rows[:3000]:
                if col_idx - 1 < len(row):
                    max_width = max(max_width, self.calc_display_width(row[col_idx - 1]))
            ws.column_dimensions[self.column_letter(col_idx)].width = min(max(max_width + 2, 10), 40)

        wb.save(path)

    def export_xlsx_minimal(self, path, headers=None, rows=None, table_name=None):
        headers = [str(h) for h in (self.headers if headers is None else headers)]
        rows = [list(row) for row in (self.rows if rows is None else rows)]
        sheet_rows = [headers]
        for row in rows:
            fixed = list(row)
            if len(fixed) < len(headers):
                fixed += [""] * (len(headers) - len(fixed))
            if len(fixed) > len(headers):
                fixed = fixed[:len(headers)]
            sheet_rows.append(["" if value is None else str(value) for value in fixed])

        def cell_xml(row_idx, col_idx, value, style_id="0"):
            ref = f"{self.column_letter(col_idx)}{row_idx}"
            value = escape(str(value))
            return f'<c r="{ref}" t="inlineStr" s="{style_id}"><is><t>{value}</t></is></c>'

        col_xml = []
        for col_idx, header in enumerate(headers, start=1):
            max_width = self.calc_display_width(header)
            for row in rows[:3000]:
                if col_idx - 1 < len(row):
                    max_width = max(max_width, self.calc_display_width(row[col_idx - 1]))
            width = min(max(max_width + 2, 10), 40)
            col_xml.append(f'<col min="{col_idx}" max="{col_idx}" width="{width}" customWidth="1"/>')

        row_xml_list = []
        for r_idx, row in enumerate(sheet_rows, start=1):
            style_id = "1" if r_idx == 1 else "0"
            cells = "".join(cell_xml(r_idx, c_idx, value, style_id) for c_idx, value in enumerate(row, start=1))
            row_xml_list.append(f'<row r="{r_idx}">{cells}</row>')

        last_col = self.column_letter(len(headers) if headers else 1)
        last_row = max(len(sheet_rows), 1)
        auto_filter = f'<autoFilter ref="A1:{last_col}{last_row}"/>' if headers else ""

        sheet_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheetViews>
    <sheetView workbookViewId="0">
      <pane ySplit="1" topLeftCell="A2" activePane="bottomLeft" state="frozen"/>
      <selection pane="bottomLeft"/>
    </sheetView>
  </sheetViews>
  <cols>{''.join(col_xml)}</cols>
  <sheetData>{''.join(row_xml_list)}</sheetData>
  {auto_filter}
</worksheet>'''

        styles_xml = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<styleSheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <fonts count="2"><font><sz val="11"/><name val="Calibri"/></font><font><b/><sz val="11"/><name val="Calibri"/></font></fonts>
  <fills count="3"><fill><patternFill patternType="none"/></fill><fill><patternFill patternType="gray125"/></fill><fill><patternFill patternType="solid"><fgColor rgb="FFD9EAF7"/></patternFill></fill></fills>
  <borders count="2"><border/><border><left style="thin"><color rgb="FFCCCCCC"/></left><right style="thin"><color rgb="FFCCCCCC"/></right><top style="thin"><color rgb="FFCCCCCC"/></top><bottom style="thin"><color rgb="FFCCCCCC"/></bottom></border></borders>
  <cellStyleXfs count="1"><xf numFmtId="0" fontId="0" fillId="0" borderId="0"/></cellStyleXfs>
  <cellXfs count="2"><xf numFmtId="0" fontId="0" fillId="0" borderId="1" xfId="0" applyBorder="1"/><xf numFmtId="0" fontId="1" fillId="2" borderId="1" xfId="0" applyFont="1" applyFill="1" applyBorder="1"/></cellXfs>
</styleSheet>'''

        content_types = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
  <Override PartName="/xl/styles.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.styles+xml"/>
  <Override PartName="/docProps/core.xml" ContentType="application/vnd.openxmlformats-package.core-properties+xml"/>
  <Override PartName="/docProps/app.xml" ContentType="application/vnd.openxmlformats-officedocument.extended-properties+xml"/>
</Types>'''

        rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties" Target="docProps/core.xml"/>
  <Relationship Id="rId3" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties" Target="docProps/app.xml"/>
</Relationships>'''

        workbook_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets><sheet name="{escape(self.normalize_sheet_title(self.table_name_var.get()))}" sheetId="1" r:id="rId1"/></sheets>
</workbook>'''

        workbook_rels = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
  <Relationship Id="rId2" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/styles" Target="styles.xml"/>
</Relationships>'''

        now = datetime.utcnow().strftime("%Y-%m-%dT%H:%M:%SZ")
        core_xml = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<cp:coreProperties xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" xmlns:dcmitype="http://purl.org/dc/dcmitype/" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
  <dc:creator>ClipboardTableTool</dc:creator><cp:lastModifiedBy>ClipboardTableTool</cp:lastModifiedBy><dcterms:created xsi:type="dcterms:W3CDTF">{now}</dcterms:created><dcterms:modified xsi:type="dcterms:W3CDTF">{now}</dcterms:modified>
</cp:coreProperties>'''

        app_xml = '''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Properties xmlns="http://schemas.openxmlformats.org/officeDocument/2006/extended-properties" xmlns:vt="http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"><Application>ClipboardTableTool</Application></Properties>'''

        with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as zf:
            zf.writestr("[Content_Types].xml", content_types)
            zf.writestr("_rels/.rels", rels)
            zf.writestr("xl/workbook.xml", workbook_xml)
            zf.writestr("xl/_rels/workbook.xml.rels", workbook_rels)
            zf.writestr("xl/worksheets/sheet1.xml", sheet_xml)
            zf.writestr("xl/styles.xml", styles_xml)
            zf.writestr("docProps/core.xml", core_xml)
            zf.writestr("docProps/app.xml", app_xml)

    def choose_db(self):
        path = filedialog.asksaveasfilename(
            title="选择 SQLite 数据库",
            defaultextension=".db",
            filetypes=[
                ("SQLite 数据库", "*.db"),
                ("SQLite 数据库", "*.sqlite"),
                ("所有文件", "*.*")
            ]
        )
        if path:
            self.db_path_var.set(path)
            self.refresh_table_list()

    def get_db_path(self):
        return self.db_path_var.get().strip()

    def get_table_names(self):
        db_path = self.get_db_path()

        if not db_path or not os.path.exists(db_path):
            return []
        return TableAccessManager(db_path, node_type="主界面").list_tables()

    def get_table_columns(self, table_name):
        db_path = self.get_db_path()

        return TableAccessManager(db_path, node_type="主界面").get_columns(table_name)

    def refresh_table_list(self):
        try:
            tables = self.get_table_names()
            self.table_combo["values"] = tables

            if tables:
                self.info_var.set(f"已读取当前数据库表：{len(tables)} 个。")
            else:
                self.info_var.set("当前数据库中没有普通数据表。")
        except Exception as e:
            self.table_combo["values"] = []
            self.info_var.set(f"读取数据库表失败：{e}")

    def sanitize_sql_name(self, name, default_name):
        return core_sanitize_sql_name(name, default_name)

    def make_sql_columns(self, headers):
        return core_make_sql_columns(headers)

    def quote_ident(self, name):
        return core_quote_ident(name)

    def table_exists(self, conn, table_name):
        cur = conn.cursor()
        cur.execute(
            "SELECT name FROM sqlite_master WHERE type='table' AND name=?",
            (table_name,)
        )
        return cur.fetchone() is not None

    def get_available_table_name(self, conn, base_name):
        if not self.table_exists(conn, base_name):
            return base_name

        suffix = datetime.now().strftime("%Y%m%d_%H%M%S")
        new_name = f"{base_name}_{suffix}"

        counter = 2
        while self.table_exists(conn, new_name):
            new_name = f"{base_name}_{suffix}_{counter}"
            counter += 1

        return new_name

    def format_db_value(self, value):
        if value is None:
            return ""

        if isinstance(value, bytes):
            return f"<BLOB {len(value)} bytes>"

        return str(value)

    def load_table_from_sqlite(self, table_name):
        db_path = self.db_path_var.get().strip()

        if not db_path:
            messagebox.showwarning("提示", "请先选择 SQLite 数据库。")
            return

        if not os.path.exists(db_path):
            messagebox.showwarning("提示", "当前数据库文件不存在。")
            return

        try:
            manager = TableAccessManager(db_path, node_type="主界面读取")
            if not manager.table_exists(table_name):
                messagebox.showwarning("提示", f"表不存在：{table_name}")
                return

            data = manager.read_table(table_name)
            headers = list(data.get("headers", []))

            if not headers:
                messagebox.showwarning("提示", f"表没有字段：{table_name}")
                return

            if self.edit_entry is not None:
                self.edit_entry.destroy()
                self.edit_entry = None

            self.raw_data = ""

            self.headers = self.make_display_headers(headers)
            self.rows = [list(row) for row in data.get("rows", [])]

            self.refresh_tree()

            self.info_var.set(
                f"已加载数据库表：{table_name}，"
                f"{len(self.rows)} 行 × {len(self.headers)} 列。"
            )

        except Exception as e:
            messagebox.showerror("读取表失败", str(e))

    def save_rows_to_sqlite_table(self, table_name_raw, headers, rows, recreate=True):
        db_path = self.get_db_path()
        if not db_path:
            raise ValueError("数据库路径为空。")

        table_name = self.sanitize_sql_name(table_name_raw, "result_table")
        sql_columns = self.make_sql_columns(headers)

        normalized_rows = []
        for row in rows:
            fixed_row = list(row)
            if len(fixed_row) < len(sql_columns):
                fixed_row += [""] * (len(sql_columns) - len(fixed_row))
            if len(fixed_row) > len(sql_columns):
                fixed_row = fixed_row[:len(sql_columns)]
            normalized_rows.append(fixed_row)

        mode = "replace" if recreate else "timestamp"
        info = TableAccessManager(db_path, node_type="主界面保存").write_table(
            table_name,
            sql_columns,
            normalized_rows,
            mode=mode,
        )
        self.refresh_table_list()

        return info.get("table_name", table_name), len(normalized_rows)

    def save_to_sqlite(self):
        if not self.headers or not self.rows:
            messagebox.showwarning("提示", "当前没有可保存的数据，请先读取剪贴板。")
            return

        db_path = self.db_path_var.get().strip()
        if not db_path:
            messagebox.showwarning("提示", "请填写 SQLite 数据库路径。")
            return

        table_name_raw = self.table_name_var.get().strip()

        try:
            table_name, row_count = self.save_rows_to_sqlite_table(
                table_name_raw=table_name_raw,
                headers=self.headers,
                rows=self.rows,
                recreate=self.recreate_table_var.get()
            )

            self.info_var.set(
                f"保存成功：数据库 {db_path}，表 {table_name}，共 {row_count} 行。"
            )

            messagebox.showinfo(
                "保存成功",
                f"已保存到 SQLite。\n\n数据库：{db_path}\n表名：{table_name}\n行数：{row_count}"
            )

        except Exception as e:
            messagebox.showerror("保存失败", str(e))

    def make_table_backup_name(self, conn, table_name):
        """生成当前 SQLite 表的备份表名，避免覆盖已有备份。"""
        base_name = f"{table_name}_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        backup_name = base_name
        counter = 2
        while self.table_exists(conn, backup_name):
            backup_name = f"{base_name}_{counter}"
            counter += 1
        return backup_name

    def backup_sqlite_table_before_delete(self, conn, table_name):
        """删除前复制当前表到同库备份表，返回备份表名。"""
        backup_name = self.make_table_backup_name(conn, table_name)
        conn.execute(
            f"CREATE TABLE {self.quote_ident(backup_name)} AS "
            f"SELECT * FROM {self.quote_ident(table_name)}"
        )
        return backup_name

    def delete_current_sqlite_table(self):
        """主页删除当前下拉框选中的 SQLite 表。

        安全规则：
        1. 必须先开启修改模式。
        2. 只允许删除当前数据库中已存在、且当前下拉框选中的普通表。
        3. 删除前询问是否备份，备份失败则不删除。
        4. 删除前进行二次确认。
        """
        if not self.edit_mode:
            messagebox.showwarning(
                "禁止删除",
                "删除 SQLite 表属于高风险操作。\n\n请先开启“修改模式:开”，再点击“删除当前表”。"
            )
            return

        db_path = self.get_db_path()
        if not db_path:
            messagebox.showwarning("提示", "请先选择 SQLite 数据库。")
            return

        if not os.path.exists(db_path):
            messagebox.showwarning("提示", "当前 SQLite 数据库文件不存在。")
            return

        table_name = self.table_name_var.get().strip()
        if not table_name:
            messagebox.showwarning("提示", "请先在“表名”下拉框选择要删除的 SQLite 表。")
            return

        try:
            tables = self.get_table_names()
        except Exception as e:
            messagebox.showerror("读取表失败", str(e))
            return

        if table_name not in tables:
            messagebox.showwarning(
                "禁止删除",
                "只能删除当前 SQLite 数据库中已存在、并且从表名下拉框选中的普通表。\n\n"
                f"当前表名：{table_name}"
            )
            return

        if table_name.lower().startswith("sqlite_"):
            messagebox.showwarning("禁止删除", "不能删除 SQLite 系统内部表。")
            return

        backup_choice = messagebox.askyesnocancel(
            "删除当前表",
            "即将删除当前 SQLite 表：\n\n"
            f"数据库：{db_path}\n"
            f"表名：{table_name}\n\n"
            "是否先备份后删除？\n\n"
            "是：先复制为备份表，再删除当前表。\n"
            "否：不备份，直接删除当前表。\n"
            "取消：放弃删除。"
        )

        if backup_choice is None:
            self.info_var.set("已取消删除当前表。")
            return

        confirm_text = (
            "请再次确认删除操作。\n\n"
            f"将删除 SQLite 表：{table_name}\n"
        )
        if backup_choice:
            confirm_text += "删除前会先在当前数据库中创建备份表。\n\n"
        else:
            confirm_text += "本次选择不备份，删除后只能依靠你自己的数据库备份恢复。\n\n"
        confirm_text += "确认继续删除吗？"

        if not messagebox.askyesno("二次确认删除", confirm_text):
            self.info_var.set("已取消删除当前表。")
            return

        try:
            backup_name = TableAccessManager(db_path, node_type="主界面删除").drop_table(
                table_name,
                backup=bool(backup_choice),
            )

            self.refresh_table_list()

            if backup_name:
                self.table_name_var.set(backup_name)
                try:
                    self.load_table_from_sqlite(backup_name)
                except Exception:
                    self.clear_preview()
                msg = f"已备份并删除当前表。备份表：{backup_name}"
            else:
                self.table_name_var.set("")
                self.clear_preview()
                msg = f"已删除当前表：{table_name}"

            self.info_var.set(msg)
            messagebox.showinfo(
                "删除完成",
                f"已删除 SQLite 表：{table_name}"
                + (f"\n\n备份表：{backup_name}" if backup_name else "\n\n本次未创建备份表。")
            )

        except Exception as e:
            messagebox.showerror("删除失败", str(e))
            self.info_var.set(f"删除失败：{e}")
